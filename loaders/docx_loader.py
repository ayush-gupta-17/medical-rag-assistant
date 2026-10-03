from pathlib import Path
from docx import Document as DocxDocument

from langchain_core.documents import Document


class DOCXLoader:

    def __init__(self, file_path):
        self.file_path = file_path

    def load(self):

        doc = DocxDocument(self.file_path)

        text = "\n".join(
            p.text
            for p in doc.paragraphs
            if p.text.strip()
        )

        return [
            Document(
                page_content=text,
                metadata={
                    "source": Path(self.file_path).name,
                    "page": 1,
                    "file_type": "docx"
                }
            )
        ]